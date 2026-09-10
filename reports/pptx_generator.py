import datetime
from pathlib import Path
from typing import Dict, Any
from pptx import Presentation
from pptx.util import Inches, Pt
from app.config import config
from tools.file_tools import get_equipment_meta

class PPTXReportGenerator:
    """Generates PowerPoint presentation summary deliverables."""

    def generate_report(self, state: Dict[str, Any]) -> Dict[str, Any]:
        prs = Presentation()
        eq_meta = get_equipment_meta(state)
        
        # Slide 1: Title
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = f"MRPL Maintenance Briefing: {eq_meta['tag']}"
        subtitle.text = (
            f"Asset: {eq_meta['desc']} ({eq_meta['unit']})\n"
            f"Risk: {state.get('risk_level', 'HIGH')} | Status: {state.get('human_approval_status', 'APPROVED')}\n"
            f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )

        output_filename = f"MRPL_{eq_meta['short']}_Briefing_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        out_path = config.output_dir / output_filename
        prs.save(out_path)

        return {"success": True, "filename": output_filename, "report_path": str(out_path)}

pptx_generator = PPTXReportGenerator()
